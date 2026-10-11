"""Czytnik .docx → lista bloków (nagłówki, akapity, listy, tabele, przypisy)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Dict, List, Optional, Union

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .settings import DocxSettings


@dataclass
class TextBlock:
    kind: str  # "heading" | "paragraph" | "list"
    text: str
    level: int = 0  # poziom nagłówka / wcięcie listy
    ordered: bool = False
    alignment: Optional[str] = None


@dataclass
class TableBlock:
    kind: str = "table"
    rows: List[List[str]] = field(default_factory=list)


@dataclass
class FootnoteBlock:
    kind: str = "footnote"
    id: str = ""
    text: str = ""


Block = Union[TextBlock, TableBlock, FootnoteBlock]

_HEADING_RE = re.compile(r"^(?:heading|nagłówek|nag\S*)\s*(\d)", re.IGNORECASE)
_LIST_STYLE_RE = re.compile(r"(list|lista|wypunkt|numer)", re.IGNORECASE)
_WS_RE = re.compile(r"[ \t\u00a0\u2007\u202f]+")


def _clean(text: str) -> str:
    text = text.replace("\r", "")
    lines = [_WS_RE.sub(" ", ln).strip() for ln in text.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def _paragraph_text(p: Paragraph, footnote_marks: bool) -> str:
    """Tekst akapitu łącznie z hiperłączami i znacznikami przypisów [^n]."""
    parts: List[str] = []
    for node in p._p.iter():
        tag = node.tag
        if tag == qn("w:t"):
            parts.append(node.text or "")
        elif tag == qn("w:tab"):
            parts.append("\t")
        elif tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
        elif tag == qn("w:noBreakHyphen"):
            parts.append("-")
        elif tag == qn("w:softHyphen"):
            continue
        elif tag == qn("w:footnoteReference") and footnote_marks:
            fid = node.get(qn("w:id"))
            if fid is not None:
                parts.append(f"[^{fid}]")
        elif tag == qn("w:delText"):
            continue
    return "".join(parts)


def _heading_level(p: Paragraph) -> int:
    try:
        name = p.style.name if p.style is not None else ""
    except Exception:  # pragma: no cover
        name = ""
    m = _HEADING_RE.match(name or "")
    if m:
        return max(1, min(6, int(m.group(1))))
    if (name or "").lower() in ("title", "tytuł"):
        return 1
    # outlineLvl w pPr
    ppr = p._p.pPr
    if ppr is not None:
        lvl = ppr.find(qn("w:outlineLvl"))
        if lvl is not None:
            try:
                return max(1, min(6, int(lvl.get(qn("w:val"))) + 1))
            except (TypeError, ValueError):
                pass
    return 0


def _is_list(p: Paragraph) -> Optional[bool]:
    """None – nie lista; True – numerowana; False – punktowana."""
    ppr = p._p.pPr
    num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
    name = ""
    try:
        name = p.style.name or ""
    except Exception:  # pragma: no cover
        pass
    if num_pr is None and not _LIST_STYLE_RE.search(name):
        return None
    ordered = bool(re.search(r"(number|numer)", name, re.IGNORECASE))
    if num_pr is not None:
        ordered = ordered or _num_fmt_is_ordered(p, num_pr)
    return ordered


def _num_fmt_is_ordered(p: Paragraph, num_pr) -> bool:
    try:
        num_id_el = num_pr.find(qn("w:numId"))
        ilvl_el = num_pr.find(qn("w:ilvl"))
        if num_id_el is None:
            return False
        num_id = num_id_el.get(qn("w:val"))
        ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"
        numbering = p.part.numbering_part.element
        num = next((n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id), None)
        if num is None:
            return False
        abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = next(
            (a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abs_id), None
        )
        if abstract is None:
            return False
        for lvl in abstract.findall(qn("w:lvl")):
            if lvl.get(qn("w:ilvl")) == ilvl:
                fmt = lvl.find(qn("w:numFmt"))
                return fmt is not None and fmt.get(qn("w:val")) not in ("bullet", "none")
    except Exception:
        return False
    return False


def _list_level(p: Paragraph) -> int:
    ppr = p._p.pPr
    if ppr is None:
        return 0
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        return 0
    ilvl = num_pr.find(qn("w:ilvl"))
    try:
        return int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
    except (TypeError, ValueError):
        return 0


def _looks_like_bold_heading(p: Paragraph, text: str) -> bool:
    if len(text) > 90 or "\n" in text or text.endswith((".", ",", ";", ":")):
        return False
    runs = [r for r in p.runs if r.text.strip()]
    if not runs:
        return False
    all_bold = all(r.bold or (r.style is not None and getattr(r.style.font, "bold", False)) for r in runs)
    if not all_bold:
        try:
            all_bold = bool(p.style is not None and p.style.font is not None and p.style.font.bold)
        except Exception:  # pragma: no cover
            all_bold = False
    centered = p.alignment is not None and int(p.alignment) == 1
    upper = text.isupper() and len(text) >= 4
    return (all_bold and (centered or upper)) or (upper and centered)


def _table_rows(tbl: Table, footnote_marks: bool) -> List[List[str]]:
    rows: List[List[str]] = []
    for row in tbl.rows:
        cells: List[str] = []
        prev = None
        for cell in row.cells:
            # scalone komórki w python-docx powtarzają ten sam obiekt
            if prev is not None and cell._tc is prev:
                cells.append("")
                continue
            prev = cell._tc
            cells.append(_clean("\n".join(_paragraph_text(p, footnote_marks) for p in cell.paragraphs)))
        rows.append(cells)
    return rows


def _footnotes(doc) -> Dict[str, str]:
    """Przypisy dolne z word/footnotes.xml (python-docx nie ma do nich API)."""
    out: Dict[str, str] = {}
    try:
        part = next(
            (rel.target_part for rel in doc.part.rels.values() if rel.reltype.endswith("/footnotes")), None
        )
        if part is None:
            return out
        from lxml import etree

        root = etree.fromstring(part.blob)
        for fn in root.iter(qn("w:footnote")):
            fid = fn.get(qn("w:id"))
            ftype = fn.get(qn("w:type"))
            if ftype in ("separator", "continuationSeparator"):
                continue
            texts: List[str] = []
            for p in fn.iter(qn("w:p")):
                parts = []
                for node in p.iter():
                    if node.tag == qn("w:t"):
                        parts.append(node.text or "")
                    elif node.tag == qn("w:tab"):
                        parts.append(" ")
                texts.append("".join(parts))
            txt = _clean(" ".join(texts))
            if fid is not None and txt:
                out[fid] = txt
    except Exception:
        return out
    return out


def read_docx(source: Union[str, bytes, BytesIO], settings: Optional[DocxSettings] = None) -> List[Block]:
    """Wczytuje dokument Word i zwraca uporządkowaną listę bloków."""
    settings = settings or DocxSettings()
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(source)
    doc = Document(source)
    blocks: List[Block] = []

    if settings.include_headers_footers:
        for section in doc.sections:
            for hp in section.header.paragraphs:
                t = _clean(_paragraph_text(hp, False))
                if t:
                    blocks.append(TextBlock("paragraph", t))

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            raw = _paragraph_text(p, settings.include_footnotes)
            text = _clean(raw)
            if not text:
                continue
            level = _heading_level(p)
            if level:
                blocks.append(TextBlock("heading", text.replace("\n", " "), level=level))
                continue
            is_list = _is_list(p)
            if is_list is not None:
                blocks.append(TextBlock("list", text, level=_list_level(p), ordered=is_list))
                continue
            if settings.detect_bold_headings and _looks_like_bold_heading(p, text):
                blocks.append(TextBlock("heading", text, level=2))
                continue
            align = None
            if p.alignment is not None:
                align = {1: "center", 2: "right"}.get(int(p.alignment))
            blocks.append(TextBlock("paragraph", text, alignment=align))
        elif child.tag == qn("w:tbl"):
            rows = _table_rows(Table(child, doc), settings.include_footnotes)
            rows = [r for r in rows if any(c for c in r)]
            if rows:
                blocks.append(TableBlock(rows=rows))
        elif child.tag == qn("w:sdt"):
            # content control – akapity w środku
            for p_el in child.iter(qn("w:p")):
                text = _clean(_paragraph_text(Paragraph(p_el, doc), settings.include_footnotes))
                if text:
                    blocks.append(TextBlock("paragraph", text))

    if settings.include_headers_footers:
        for section in doc.sections:
            for fp in section.footer.paragraphs:
                t = _clean(_paragraph_text(fp, False))
                if t:
                    blocks.append(TextBlock("paragraph", t))

    if settings.include_footnotes:
        for fid, text in _footnotes(doc).items():
            blocks.append(FootnoteBlock(id=fid, text=text))

    return blocks
